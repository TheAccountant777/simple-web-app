import io
from typing import Any

import openpyxl
import pytest
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle

from kenya_data_engine.config import DataConfig
from kenya_data_engine.data.extract import Locator, RawTable, extract_tables, table_to_frame
from kenya_data_engine.errors import ExtractError

CFG = DataConfig(max_bytes={"html": 1_000_000, "pdf": 1_000_000, "sheet": 1_000_000})


def _xlsx(build: Any) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sheet1"
    build(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _merged(ws: Any) -> None:
    ws["B1"] = "Pump prices"
    ws.merge_cells("B1:C1")
    ws["A2"], ws["B2"], ws["C2"] = "Town", "Super", "Diesel"
    ws.append(["Nairobi", 180.66, 170.5])
    ws.append(["Mombasa", 179.1, 169.0])


async def test_xlsx_merged_multirow_header() -> None:
    (t,) = await extract_tables(_xlsx(_merged), Locator(), CFG)
    assert t.header == ["Town", "Pump prices / Super", "Pump prices / Diesel"]
    assert t.rows[0] == ["Nairobi", "180.66", "170.5"]
    assert t.cell_locators[0][1] == "Sheet1!B3"
    assert t.source_kind == "xlsx" and t.extractor.startswith("openpyxl@")


async def test_footnote_rows_dropped() -> None:
    def build(ws: Any) -> None:
        ws.append(["Town", "Price"])
        ws.append(["Nairobi", 1])
        ws.append(["Source: EPRA", None])
        ws.append(["* provisional"])
        ws.append(["1) note"])

    (t,) = await extract_tables(_xlsx(build), Locator(), CFG)
    assert t.rows == [["Nairobi", "1"]]


async def test_blank_leading_column_keeps_letters() -> None:
    def build(ws: Any) -> None:
        ws.append([None, "Town", "Price"])
        ws.append([None, "Nairobi", 5])

    (t,) = await extract_tables(_xlsx(build), Locator(), CFG)
    assert t.header == ["Town", "Price"]
    assert t.cell_locators[0] == ["Sheet1!B2", "Sheet1!C2"]


async def test_xlsx_sheet_selection_and_errors() -> None:
    def build(ws: Any) -> None:
        ws.append(["A"])
        ws.append([1])

    data = _xlsx(build)
    assert len(await extract_tables(data, Locator(sheet="Sheet1"), CFG)) == 1
    with pytest.raises(ExtractError):
        await extract_tables(data, Locator(sheet="Nope"), CFG)


async def test_explicit_header_rows_zero_and_frame() -> None:
    csv = b"a,b\n1,2\n3,4\n"
    (t,) = await extract_tables(csv, Locator(header_rows=0), CFG)
    assert t.rows[0] == ["a", "b"] and t.header == ["", ""]
    (t2,) = await extract_tables(csv, Locator(header_rows=1), CFG)
    df = table_to_frame(t2)
    assert list(df.columns) == ["a", "b"] and df.iloc[1, 1] == "4"


async def test_csv_thousands_text_cells() -> None:
    csv = b'Town,Value\n"Nairobi","1,234.5"\nMombasa,99\n'
    (t,) = await extract_tables(csv, Locator(), CFG)
    assert t.header == ["Town", "Value"]
    assert t.rows[0] == ["Nairobi", "1,234.5"]
    assert t.cell_locators[0][1] == "t0/r1/c1"
    semi = b"Town;Value\nNairobi;1\n"
    (t2,) = await extract_tables(semi, Locator(), CFG)
    assert t2.rows == [["Nairobi", "1"]]


async def test_csv_latin1_fallback() -> None:
    (t,) = await extract_tables("Town,V\nNairobi \xe9,1\n".encode("latin-1"), Locator(), CFG)
    assert t.rows[0][0] == "Nairobi \xe9"


async def test_html_table_by_css() -> None:
    html = (
        "<html><body><table id='a'><tr><th>X</th></tr><tr><td>1</td></tr></table>"
        "<table id='prices'><tr><th>Town</th><th colspan='2'>Pump</th></tr>"
        "<tr><td>Nairobi</td><td>180.66</td><td>170.5</td></tr></table>"
        "<div id='wrap'><table><tr><td>Q</td></tr></table></div></body></html>"
    )
    (t,) = await extract_tables(html.encode(), Locator(css="table#prices"), CFG)
    assert t.header == ["Town", "Pump", "Pump"]
    assert t.rows == [["Nairobi", "180.66", "170.5"]]
    assert t.cell_locators[0][2] == "t0/r1/c2"
    assert t.source_kind == "html"
    assert len(await extract_tables(html.encode(), Locator(), CFG)) == 3
    wrapped = await extract_tables(html.encode(), Locator(css="#wrap"), CFG)
    assert wrapped[0].header == ["Q"] and wrapped[0].rows == []


def _pdf(pages: int = 2) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    style = TableStyle([("GRID", (0, 0), (-1, -1), 1, (0, 0, 0))])
    story: list[Any] = []
    for p in range(1, pages + 1):
        story.append(Paragraph(f"Page {p}", getSampleStyleSheet()["Normal"]))
        if p == 2:
            tbl = Table([["Town", "Super"], ["Nairobi", "180.66"], ["Mombasa", "179.10"]])
            tbl.setStyle(style)
            story.append(tbl)
        if p < pages:
            story.append(PageBreak())
    doc.build(story)
    return buf.getvalue()


async def test_pdf_table_page_locator() -> None:
    (t,) = await extract_tables(_pdf(), Locator(pages=[2]), CFG)
    assert t.header == ["Town", "Super"]
    assert t.rows[0] == ["Nairobi", "180.66"]
    assert t.cell_locators[0][0].startswith("p2/")
    assert t.cell_locators[0][1] == "p2/t0/r1/c1"
    assert t.extractor.startswith("pdfplumber@") and t.source_kind == "pdf"
    assert len(await extract_tables(_pdf(), Locator(), CFG)) == 1


async def test_pdf_page_cap_raises() -> None:
    cfg = CFG.model_copy(update={"pdf_max_pages": 1})
    with pytest.raises(ExtractError):
        await extract_tables(_pdf(), Locator(), cfg)
    with pytest.raises(ExtractError):
        await extract_tables(_pdf(), Locator(pages=[9]), CFG)


async def test_pdf_garbage_raises() -> None:
    with pytest.raises(ExtractError):
        await extract_tables(b"%PDF-1.4 not really", Locator(), CFG)


async def test_xlsx_zip_bomb_guard() -> None:
    def build(ws: Any) -> None:
        for _ in range(300):
            ws.append(["a" * 1000])

    data = _xlsx(build)
    cfg = CFG.model_copy(update={"max_bytes": {"html": 1, "pdf": 1, "sheet": len(data) + 1}})
    with pytest.raises(ExtractError, match="expands"):
        await extract_tables(data, Locator(), cfg)


async def test_unknown_type_raises() -> None:
    with pytest.raises(ExtractError):
        await extract_tables(b"\x00\x01\x02", Locator(), CFG)
    with pytest.raises(ExtractError):
        await extract_tables(b'{"a": 1}', Locator(), CFG)


async def test_legacy_xls_unsupported() -> None:
    with pytest.raises(ExtractError, match="xls"):
        await extract_tables(bytes.fromhex("D0CF11E0") + b"\x00" * 20, Locator(), CFG)


async def test_timeout_raises(monkeypatch: pytest.MonkeyPatch) -> None:
    import time

    from kenya_data_engine.data import extract

    monkeypatch.setattr(extract, "_dispatch", lambda *a: time.sleep(0.5))
    cfg = CFG.model_copy(update={"parse_timeout_s": 0.05})
    with pytest.raises(ExtractError, match="longer"):
        await extract.extract_tables(b"a,b", Locator(), cfg)


def test_rawtable_model() -> None:
    t = RawTable(
        header=["a"], rows=[["1"]], cell_locators=[["t0/r1/c0"]], extractor="x", source_kind="csv"
    )
    assert table_to_frame(t).iloc[0, 0] == "1"


async def test_unmerged_blank_upper_cell_stays_blank() -> None:
    def build(ws: Any) -> None:
        ws.append(["Town", None, "Pump", None])
        ws.append([None, "Region", "Super", "Diesel"])
        ws.append(["Nairobi", "Central", 180.5, 170.5])

    (t,) = await extract_tables(_xlsx(build), Locator(header_rows=2), CFG)
    assert t.header == ["Town", "Region", "Pump / Super", "Diesel"]


async def test_merged_range_fills_header() -> None:
    (t,) = await extract_tables(_xlsx(_merged), Locator(), CFG)
    assert t.header[1:] == ["Pump prices / Super", "Pump prices / Diesel"]


async def test_html_colspan_rowspan_align() -> None:
    html = (
        "<table><tr><th rowspan='2'>Town</th><th colspan='2'>Pump</th></tr>"
        "<tr><th>Super</th><th>Diesel</th></tr>"
        "<tr><td>Nairobi</td><td>180</td><td>170</td></tr></table>"
    )
    (t,) = await extract_tables(html.encode(), Locator(header_rows=2), CFG)
    assert t.header == ["Town", "Pump / Super", "Pump / Diesel"]
    assert t.rows == [["Nairobi", "180", "170"]]


async def test_notional_row_not_footnote() -> None:
    csv = b"Town,Value\nNotional,1\nSourcing,2\nSource: EPRA,\nNote 1,\n"
    (t,) = await extract_tables(csv, Locator(), CFG)
    assert t.rows == [["Notional", "1"], ["Sourcing", "2"]]


async def test_corrupt_xlsx_raises_extract_error() -> None:
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("xl/workbook.xml", "<not xml")
    with pytest.raises(ExtractError):
        await extract_tables(buf.getvalue(), Locator(), CFG)


async def test_html_spanned_copy_is_marked_except_label_column() -> None:
    html = (
        b"<table><tr><th>Town</th><th>Super</th><th>Diesel</th></tr>"
        b"<tr><td>Nairobi</td><td colspan=2>180.50</td></tr>"
        b"<tr><td rowspan=2>Coast</td><td>1</td><td>2</td></tr>"
        b"<tr><td>3</td><td>4</td></tr></table>"
    )
    (t,) = await extract_tables(html, Locator(), CFG)
    assert t.rows[0] == ["Nairobi", "180.50", "180.50"]
    assert t.cell_locators[0] == ["t0/r1/c0", "t0/r1/c1", "span:t0/r1/c1"]
    assert t.cell_locators[2][0] == "t0/r2/c0"  # label column: the copy is a label, not a number


async def test_xlsx_merged_copy_is_marked() -> None:
    def build(ws: Any) -> None:
        ws.append(["Town", "Super", "Diesel"])
        ws.append(["Nairobi", 180.5, None])
        ws.merge_cells("B2:C2")

    (t,) = await extract_tables(_xlsx(build), Locator(), CFG)
    assert t.rows[0] == ["Nairobi", "180.5", "180.5"]
    assert t.cell_locators[0][1:] == ["Sheet1!B2", "span:Sheet1!B2"]


async def test_star_followed_by_letter_is_not_a_footnote() -> None:
    def build(ws: Any) -> None:
        ws.append(["Town", "Price"])
        ws.append(["*Nairobi", 1])
        ws.append(["*2 provisional", None])
        ws.append(["* provisional", None])

    (t,) = await extract_tables(_xlsx(build), Locator(), CFG)
    assert t.rows == [["*Nairobi", "1"]]


async def test_sheet_size_cap_enforced_for_xlsx_and_csv() -> None:
    small = DataConfig(max_bytes={"html": 1000, "pdf": 1000, "sheet": 50})
    with pytest.raises(ExtractError, match="too large"):
        await extract_tables(b"a,b\n" + b"1,2\n" * 50, Locator(), small)

    def build(ws: Any) -> None:
        ws.append(["a"] * 40)

    with pytest.raises(ExtractError, match="too large"):
        await extract_tables(_xlsx(build), Locator(), small)
