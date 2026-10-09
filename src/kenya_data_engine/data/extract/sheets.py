"""CSV and XLSX tables."""

import csv
import io
import zipfile
from datetime import date, datetime

import openpyxl
from openpyxl.utils import get_column_letter

from kenya_data_engine.data.extract.grid import Locator, RawTable, build_table, clean
from kenya_data_engine.errors import ExtractError

ZIP_RATIO = 10


def _text(value: object) -> str:
    if isinstance(value, datetime | date):
        return value.isoformat()
    return clean(value)


def read_xlsx(content: bytes, locator: Locator, max_bytes: int) -> list[RawTable]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            if sum(i.file_size for i in zf.infolist()) > ZIP_RATIO * max_bytes:
                raise ExtractError("xlsx expands beyond the allowed size", hint="Zip bomb guard.")
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ExtractError(f"unreadable xlsx: {exc}") from exc
    try:
        name = locator.sheet or wb.sheetnames[0]
        if name not in wb.sheetnames:
            raise ExtractError(f"sheet {name!r} not found", hint=f"Sheets: {wb.sheetnames}")
        cells: list[list[str]] = []
        locs: list[list[str]] = []
        for r, row in enumerate(wb[name].iter_rows(values_only=True), start=1):
            cells.append([_text(v) for v in row])
            locs.append([f"{name}!{get_column_letter(c)}{r}" for c in range(1, len(row) + 1)])
    finally:
        wb.close()
    table = build_table(cells, locs, locator, f"openpyxl@{openpyxl.__version__}", "xlsx")
    return [table] if table else []


def read_csv(content: bytes, locator: Locator) -> list[RawTable]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(text[:4096], ",;\t|")
    except csv.Error:
        dialect = csv.excel
    cells = [[clean(c) for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    locs = [[f"t0/r{r}/c{c}" for c in range(len(row))] for r, row in enumerate(cells)]
    table = build_table(cells, locs, locator, "csv@stdlib", "csv")
    return [table] if table else []
